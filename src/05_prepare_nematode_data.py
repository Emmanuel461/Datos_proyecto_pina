"""Prepara los datos de nematodos sin modificar el libro original.

Las cuatro hojas cumplen funciones distintas:
- Nematodos en piña: conteos, parámetros y cálculos de abundancia.
- NINJA: matriz de abundancias enviada al calculador NINJA.
- INDICES: resultados devueltos por NINJA.
- NEMATODES: rasgos asignados a cada taxón.

El script conserva las cuatro capas, valida su correspondencia y genera tablas
separadas en processed/nematodos_analisis_listo. No genera archivos para la web.
"""
from __future__ import annotations

import csv
import hashlib
import math
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "raw" / "nematodos" / "nematodos_raw.xlsx"
DESIGN = ROOT / "processed" / "diseno_muestreo.csv"
OUT = ROOT / "processed" / "nematodos_analisis_listo"

COMMUNITY_SHEET = "Nematodos en piña"
INDEX_SHEET = "INDICES"
NINJA_SHEET = "NINJA"
TRAIT_SHEET = "NEMATODES"
REQUIRED_SHEETS = {COMMUNITY_SHEET, INDEX_SHEET, NINJA_SHEET, TRAIT_SHEET}

FIRST_SAMPLE_COLUMN = 2
LAST_SAMPLE_COLUMN = 13
FIRST_TAXON_ROW = 6
LAST_TAXON_ROW = 40
NINJA_FIRST_TAXON_COLUMN = 3
NINJA_LAST_TAXON_COLUMN = 37


def clean_text(value: object) -> str:
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def ascii_text(value: object) -> str:
    return unicodedata.normalize("NFKD", clean_text(value)).encode("ascii", "ignore").decode()


def slug(value: object) -> str:
    text = ascii_text(value).lower().replace("%", " pct ").replace("+", " plus ")
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def as_number(value: object, context: str, *, positive: bool = False) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{context}: se esperaba un número y se encontró {value!r}.")
    result = float(value)
    if not math.isfinite(result) or (positive and result <= 0):
        condition = "positivo" if positive else "finito"
        raise ValueError(f"{context}: se esperaba un número {condition} y se encontró {value!r}.")
    return result


def normalize_sample(value: object) -> tuple[str, str, str]:
    compact = re.sub(r"[^A-Z0-9]+", "", ascii_text(value).upper())
    match = re.fullmatch(r"(PINA|BOSQUE)(\d+)", compact)
    if not match:
        raise ValueError(f"Identificador de parcela no reconocido: {value!r}.")
    treatment = "Pina" if match.group(1) == "PINA" else "Bosque"
    lot = str(int(match.group(2)))
    return treatment, lot, f"{treatment}_{lot}"


def read_design(path: Path) -> tuple[dict[tuple[str, str], dict[str, str]], dict[int, str]]:
    if not path.exists():
        raise FileNotFoundError(
            f"No existe {path}. Ejecuta primero src/01_soil_data_pipeline.py."
        )
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    required = {"finca_id", "tratamiento", "lote", "parcela", "productor_ficha"}
    if not rows or not required.issubset(rows[0]):
        raise ValueError(f"{path.name} no contiene las columnas requeridas: {sorted(required)}.")

    design: dict[tuple[str, str], dict[str, str]] = {}
    producer_by_farm: dict[int, str] = {}
    for row in rows:
        key = (clean_text(row["tratamiento"]), str(int(row["lote"])))
        if key in design:
            raise ValueError(f"La ficha contiene más de una fila para {key}.")
        farm = int(row["finca_id"])
        design[key] = row
        producer = clean_text(row["productor_ficha"])
        if producer:
            previous = producer_by_farm.setdefault(farm, producer)
            if previous != producer:
                raise ValueError(f"La finca {farm} tiene productores distintos en la ficha.")
    return design, producer_by_farm


def sample_metadata(
    original_id: object,
    design: dict[tuple[str, str], dict[str, str]],
    producer_by_farm: dict[int, str],
) -> dict[str, object]:
    treatment, lot, parcel = normalize_sample(original_id)
    try:
        design_row = design[(treatment, lot)]
    except KeyError as error:
        raise ValueError(f"{original_id!r} no tiene correspondencia en el diseño de muestreo.") from error
    if clean_text(design_row["parcela"]) != parcel:
        raise ValueError(f"La parcela normalizada {parcel} no coincide con la ficha de campo.")
    farm = int(design_row["finca_id"])
    return {
        "finca_id": farm,
        "tratamiento": treatment,
        "lote": lot,
        "parcela": parcel,
        "productor_finca": producer_by_farm.get(farm, ""),
        "muestra_origen": clean_text(original_id),
    }


def write_csv(path: Path, columns: list[str], rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    try:
        with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def prepare() -> dict[str, tuple[list[str], list[dict[str, object]]]]:
    if not SOURCE.exists():
        raise FileNotFoundError(f"No existe el libro de nematodos: {SOURCE}")
    design, producer_by_farm = read_design(DESIGN)
    try:
        workbook = load_workbook(SOURCE, data_only=False, read_only=True)
    except PermissionError as error:
        raise PermissionError(
            f"No se puede leer {SOURCE.name}. Cierra el archivo en Excel y vuelve a ejecutar el código."
        ) from error
    try:
        if set(workbook.sheetnames) != REQUIRED_SHEETS:
            missing = sorted(REQUIRED_SHEETS - set(workbook.sheetnames))
            extra = sorted(set(workbook.sheetnames) - REQUIRED_SHEETS)
            raise ValueError(f"Hojas inesperadas. Faltantes: {missing}; adicionales: {extra}.")

        community = workbook[COMMUNITY_SHEET]
        ninja = workbook[NINJA_SHEET]
        indices = workbook[INDEX_SHEET]
        traits = workbook[TRAIT_SHEET]

        samples: dict[str, dict[str, object]] = {}
        for column in range(FIRST_SAMPLE_COLUMN, LAST_SAMPLE_COLUMN + 1):
            original_id = community.cell(2, column).value
            metadata = sample_metadata(original_id, design, producer_by_farm)
            key = metadata["parcela"]
            if key in samples:
                raise ValueError(f"Parcela duplicada en {COMMUNITY_SHEET}: {key}.")
            metadata["columna_comunidad"] = column
            samples[str(key)] = metadata
        if len(samples) != 12:
            raise ValueError(f"Se esperaban 12 parcelas y se encontraron {len(samples)}.")

        raw_taxa: list[tuple[int, str]] = []
        for row in range(FIRST_TAXON_ROW, LAST_TAXON_ROW + 1):
            taxon = clean_text(community.cell(row, 1).value)
            if not taxon:
                raise ValueError(f"Taxón vacío en {COMMUNITY_SHEET}!A{row}.")
            raw_taxa.append((row, taxon))
        if len({taxon for _, taxon in raw_taxa}) != len(raw_taxa):
            raise ValueError("Hay taxones duplicados en la tabla de conteos.")

        ninja_taxa: dict[str, int] = {}
        for column in range(NINJA_FIRST_TAXON_COLUMN, NINJA_LAST_TAXON_COLUMN + 1):
            taxon = clean_text(ninja.cell(1, column).value)
            if not taxon or taxon in ninja_taxa:
                raise ValueError(f"Taxón vacío o duplicado en {NINJA_SHEET}, columna {column}.")
            ninja_taxa[taxon] = column
        raw_taxon_names = {taxon for _, taxon in raw_taxa}
        if raw_taxon_names != set(ninja_taxa):
            raise ValueError("Los taxones de la hoja de conteos no coinciden con la matriz NINJA.")

        ninja_samples: dict[str, int] = {}
        ninja_order: dict[str, object] = {}
        for row in range(2, 14):
            metadata = sample_metadata(ninja.cell(row, 2).value, design, producer_by_farm)
            key = str(metadata["parcela"])
            if key in ninja_samples:
                raise ValueError(f"Parcela duplicada en {NINJA_SHEET}: {key}.")
            ninja_samples[key] = row
            ninja_order[key] = ninja.cell(row, 1).value
        if set(ninja_samples) != set(samples):
            raise ValueError("Las parcelas de NINJA no coinciden con las de la hoja de conteos.")

        abundance_rows: list[dict[str, object]] = []
        parameter_rows: list[dict[str, object]] = []
        matrix_rows: list[dict[str, object]] = []
        blank_counts: list[str] = []
        max_abundance_difference = 0.0

        for parcel, metadata in samples.items():
            column = int(metadata["columna_comunidad"])
            ninja_row = ninja_samples[parcel]
            abundance_1ml = as_number(
                community.cell(3, column).value,
                f"{COMMUNITY_SHEET}!{get_column_letter(column)}3",
                positive=True,
            )
            counted_volume = as_number(
                community.cell(4, column).value,
                f"{COMMUNITY_SHEET}!{get_column_letter(column)}4",
                positive=True,
            )
            id_volume = as_number(
                community.cell(5, column).value,
                f"{COMMUNITY_SHEET}!{get_column_letter(column)}5",
                positive=True,
            )

            counts: dict[str, float | None] = {}
            for row, taxon in raw_taxa:
                value = community.cell(row, column).value
                if value is None:
                    counts[taxon] = None
                    blank_counts.append(f"{parcel}:{taxon}:{get_column_letter(column)}{row}")
                else:
                    count = as_number(value, f"Conteo {parcel}, {taxon}")
                    if count < 0:
                        raise ValueError(f"Conteo negativo en {parcel}, {taxon}.")
                    counts[taxon] = count
            total_identified = sum(value for value in counts.values() if value is not None)
            if total_identified <= 0:
                raise ValueError(f"El total identificado de {parcel} no es positivo.")

            ninja_values: dict[str, float] = {}
            for _, taxon in raw_taxa:
                value = as_number(
                    ninja.cell(ninja_row, ninja_taxa[taxon]).value,
                    f"{NINJA_SHEET}, {parcel}, {taxon}",
                )
                if value < 0:
                    raise ValueError(f"Abundancia negativa en {parcel}, {taxon}.")
                ninja_values[taxon] = value
            ninja_total = sum(ninja_values.values())
            if ninja_total <= 0:
                raise ValueError(f"La abundancia total NINJA de {parcel} no es positiva.")

            parameter_rows.append({
                **{key: metadata[key] for key in (
                    "finca_id", "tratamiento", "lote", "parcela", "productor_finca", "muestra_origen"
                )},
                "numero_muestra_ninja": ninja_order[parcel],
                "abundancia_en_1_ml": abundance_1ml,
                "volumen_contado_abundancia_ml": counted_volume,
                "volumen_agregado_identificacion_ml": id_volume,
                "total_individuos_identificados": total_identified,
                "abundancia_total_ninja": ninja_total,
                "unidad_abundancia": "individuos/100 cm3 de suelo (según encabezado de la fuente)",
                "archivo_origen": SOURCE.name,
                "hoja_origen": COMMUNITY_SHEET,
            })

            matrix_row = {
                **{key: metadata[key] for key in (
                    "finca_id", "tratamiento", "lote", "parcela", "productor_finca", "muestra_origen"
                )},
                "numero_muestra_ninja": ninja_order[parcel],
            }
            matrix_row.update(ninja_values)
            matrix_rows.append(matrix_row)

            for raw_row, taxon in raw_taxa:
                raw_count = counts[taxon]
                count_used_by_source = 0.0 if raw_count is None else raw_count
                calculated = (
                    (count_used_by_source / total_identified * (abundance_1ml * id_volume))
                    * (10 / id_volume)
                )
                ninja_value = ninja_values[taxon]
                difference = abs(calculated - ninja_value)
                max_abundance_difference = max(max_abundance_difference, difference)
                if not math.isclose(calculated, ninja_value, rel_tol=1e-9, abs_tol=1e-9):
                    raise ValueError(
                        f"La abundancia de {taxon} en {parcel} no coincide con la matriz NINJA: "
                        f"{calculated} frente a {ninja_value}."
                    )
                abundance_rows.append({
                    **{key: metadata[key] for key in (
                        "finca_id", "tratamiento", "lote", "parcela", "productor_finca", "muestra_origen"
                    )},
                    "taxon": taxon,
                    "conteo_identificacion": "" if raw_count is None else raw_count,
                    "conteo_usado_por_fuente": count_used_by_source,
                    "estado_conteo": (
                        "no_reportado_tratado_como_cero_en_fuente" if raw_count is None else "reportado"
                    ),
                    "abundancia_absoluta": ninja_value,
                    "abundancia_relativa": ninja_value / ninja_total,
                    "unidad_abundancia": "individuos/100 cm3 de suelo (según encabezado de la fuente)",
                    "archivo_origen": SOURCE.name,
                    "hoja_conteo": COMMUNITY_SHEET,
                    "celda_conteo": f"{get_column_letter(column)}{raw_row}",
                    "hoja_abundancia": NINJA_SHEET,
                    "celda_abundancia": f"{get_column_letter(ninja_taxa[taxon])}{ninja_row}",
                })

        trait_rows: list[dict[str, object]] = []
        for row in range(2, traits.max_row + 1):
            taxon = clean_text(traits.cell(row, 1).value)
            if not taxon:
                continue
            trait_rows.append({
                "taxon": taxon,
                "clase_cp": traits.cell(row, 2).value,
                "clase_pp": traits.cell(row, 3).value,
                "grupo_alimenticio": clean_text(traits.cell(row, 4).value),
                "masa_ug": as_number(traits.cell(row, 5).value, f"{TRAIT_SHEET}!E{row}"),
                "archivo_origen": SOURCE.name,
                "hoja_origen": TRAIT_SHEET,
                "fila_origen": row,
            })
        if len({row["taxon"] for row in trait_rows}) != len(trait_rows):
            raise ValueError("Hay taxones duplicados en el catálogo NEMATODES.")
        if {row["taxon"] for row in trait_rows} != raw_taxon_names:
            raise ValueError("El catálogo de rasgos no contiene exactamente los taxones observados.")

        index_headers = [clean_text(indices.cell(1, column).value) for column in range(2, indices.max_column + 1)]
        index_codes = [slug(header) for header in index_headers]
        if any(not value for value in index_codes) or len(set(index_codes)) != len(index_codes):
            raise ValueError("Los nombres de índices no producen identificadores únicos.")
        index_samples: dict[str, int] = {}
        for row in range(2, indices.max_row + 1):
            metadata = sample_metadata(indices.cell(row, 1).value, design, producer_by_farm)
            parcel = str(metadata["parcela"])
            if parcel in index_samples:
                raise ValueError(f"Parcela duplicada en {INDEX_SHEET}: {parcel}.")
            index_samples[parcel] = row
        if set(index_samples) != set(samples):
            raise ValueError("Las parcelas de INDICES no coinciden con las demás hojas.")

        index_rows: list[dict[str, object]] = []
        index_wide_rows: list[dict[str, object]] = []
        max_total_difference = 0.0
        total_index_position = index_headers.index("Total number, ind") + 2
        for parcel, metadata in samples.items():
            row = index_samples[parcel]
            wide = {key: metadata[key] for key in (
                "finca_id", "tratamiento", "lote", "parcela", "productor_finca", "muestra_origen"
            )}
            for offset, (original_name, code) in enumerate(zip(index_headers, index_codes), start=2):
                value = as_number(indices.cell(row, offset).value, f"{INDEX_SHEET}!{get_column_letter(offset)}{row}")
                wide[code] = value
                index_rows.append({
                    **{key: metadata[key] for key in (
                        "finca_id", "tratamiento", "lote", "parcela", "productor_finca", "muestra_origen"
                    )},
                    "indice": code,
                    "nombre_indice_origen": original_name,
                    "valor": value,
                    "archivo_origen": SOURCE.name,
                    "hoja_origen": INDEX_SHEET,
                    "celda_origen": f"{get_column_letter(offset)}{row}",
                })
            index_wide_rows.append(wide)
            total_index = as_number(
                indices.cell(row, total_index_position).value,
                f"{INDEX_SHEET}!{get_column_letter(total_index_position)}{row}",
            )
            ninja_total = next(item["abundancia_total_ninja"] for item in parameter_rows if item["parcela"] == parcel)
            difference = abs(total_index - float(ninja_total))
            max_total_difference = max(max_total_difference, difference)
            if not math.isclose(total_index, float(ninja_total), rel_tol=1e-9, abs_tol=1e-9):
                raise ValueError(f"El total de INDICES no coincide con NINJA para {parcel}.")

        audit_rows = [
            {"control": "archivo_origen", "resultado": SOURCE.name, "detalle": "Libro preservado sin cambios"},
            {"control": "sha256_origen", "resultado": hashlib.sha256(SOURCE.read_bytes()).hexdigest(), "detalle": "Huella del libro procesado"},
            {"control": "hojas_revisadas", "resultado": len(REQUIRED_SHEETS), "detalle": " | ".join(sorted(REQUIRED_SHEETS))},
            {"control": "parcelas", "resultado": len(samples), "detalle": "8 Piña y 4 Bosque"},
            {"control": "taxones", "resultado": len(raw_taxa), "detalle": "Coinciden en conteos, NINJA y NEMATODES"},
            {"control": "indices_ninja", "resultado": len(index_headers), "detalle": "Sin valores vacíos"},
            {"control": "conteos_no_reportados", "resultado": len(blank_counts), "detalle": " | ".join(blank_counts) or "Ninguno"},
            {"control": "tratamiento_fuente_conteos_vacios", "resultado": "cero", "detalle": "Se conserva el vacío y se registra el cero usado por la fuente"},
            {"control": "diferencia_max_abundancia_vs_ninja", "resultado": max_abundance_difference, "detalle": "Tolerancia 1e-9"},
            {"control": "diferencia_max_total_indices_vs_ninja", "resultado": max_total_difference, "detalle": "Tolerancia 1e-9"},
            {"control": "volumen_contado_en_formula", "resultado": "no", "detalle": "Se conserva como parámetro; la fórmula del libro no referencia la fila 4"},
            {"control": "version_ninja", "resultado": "no reportada", "detalle": "El libro no identifica la versión del calculador"},
        ]

        identity_columns = [
            "finca_id", "tratamiento", "lote", "parcela", "productor_finca", "muestra_origen"
        ]
        return {
            "abundancia_taxones.csv": (
                identity_columns + [
                    "taxon", "conteo_identificacion", "conteo_usado_por_fuente", "estado_conteo",
                    "abundancia_absoluta", "abundancia_relativa", "unidad_abundancia", "archivo_origen",
                    "hoja_conteo", "celda_conteo", "hoja_abundancia", "celda_abundancia",
                ],
                abundance_rows,
            ),
            "parametros_muestras.csv": (
                identity_columns + [
                    "numero_muestra_ninja", "abundancia_en_1_ml", "volumen_contado_abundancia_ml",
                    "volumen_agregado_identificacion_ml", "total_individuos_identificados",
                    "abundancia_total_ninja", "unidad_abundancia", "archivo_origen", "hoja_origen",
                ],
                parameter_rows,
            ),
            "rasgos_taxones.csv": (
                ["taxon", "clase_cp", "clase_pp", "grupo_alimenticio", "masa_ug", "archivo_origen", "hoja_origen", "fila_origen"],
                trait_rows,
            ),
            "indices_ninja.csv": (
                identity_columns + ["indice", "nombre_indice_origen", "valor", "archivo_origen", "hoja_origen", "celda_origen"],
                index_rows,
            ),
            "indices_ninja_por_parcela.csv": (identity_columns + index_codes, index_wide_rows),
            "trazabilidad/matriz_ninja.csv": (
                identity_columns + ["numero_muestra_ninja"] + [taxon for _, taxon in raw_taxa],
                matrix_rows,
            ),
            "auditoria_preparacion.csv": (["control", "resultado", "detalle"], audit_rows),
        }
    finally:
        workbook.close()


def main() -> None:
    outputs = prepare()
    for relative_path, (columns, rows) in outputs.items():
        path = OUT / relative_path
        write_csv(path, columns, rows)
        print(f" - {path.relative_to(ROOT)}: {len(rows)} filas")


if __name__ == "__main__":
    main()
