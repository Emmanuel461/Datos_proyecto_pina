"""Prepara el reporte químico preservando resultados censurados y su procedencia."""
from __future__ import annotations

import csv
from pathlib import Path
import math
import re

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'processed/quimico_all_pina/quimico_all.xlsx'
OUT = ROOT / 'processed/quimico_analisis_listo'
VARIABLES = [
    ('ph', 'H₂O', 'adimensional'), ('acidez', 'Acidez', 'cmol(+)/L'),
    ('ca', 'Ca', 'cmol(+)/L'), ('mg', 'Mg', 'cmol(+)/L'),
    ('k', 'K', 'cmol(+)/L'), ('cice', 'CICE', 'cmol(+)/L'),
    ('sa', 'SA', '%'), ('p', 'P', 'mg/L'), ('zn', 'Zn', 'mg/L'),
    ('cu', 'Cu', 'mg/L'), ('fe', 'Fe', 'mg/L'), ('mn', 'Mn', 'mg/L'),
    ('ce', 'CE', 'mS/cm'), ('c', 'C', '%'), ('n', 'N', '%'),
    ('c_n', 'C/N', 'adimensional'),
]
PUBLIC_COLUMNS = [
    'parcela', 'tratamiento', 'lote', 'finca_id', 'productor_finca', 'variable', 'unidad',
    'resultado_original', 'valor', 'operador', 'limite_reportado', 'estado',
]


def parse_result(value):
    """Un límite nunca se almacena como una concentración medida."""
    if value is None or str(value).strip() == '':
        return None, '', None, 'faltante'
    match = re.fullmatch(r'\s*(<=|>=|<|>|≤|≥)?\s*([+-]?\d+(?:[.,]\d+)?(?:[eE][+-]?\d+)?)\s*', str(value))
    if not match:
        raise ValueError(f'Resultado no reconocido: {value!r}')
    operator = (match[1] or '=').replace('≤', '<=').replace('≥', '>=')
    number = float(match[2].replace(',', '.'))
    if not math.isfinite(number) or number < 0:
        raise ValueError(f'Resultado inválido: {value!r}')
    if operator == '=':
        return number, operator, None, 'cuantificado'
    return None, operator, number, 'censurado'


def prepare_dashboard_data(results: list[dict]) -> list[dict]:
    """Valida la correspondencia antes de sustituir cualquiera de las salidas."""
    design_path = ROOT / 'processed' / 'diseno_muestreo.csv'
    required_design = {'tratamiento', 'lote', 'finca_id', 'productor_ficha'}
    design = {}
    producer_by_farm = {}
    with design_path.open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle)
        if required_design - set(reader.fieldnames or []):
            raise ValueError(f'Faltan columnas de diseño en {design_path}')
        for row in reader:
            identifiers = {}
            for column in ['lote', 'finca_id']:
                try:
                    value = float(row[column])
                except (TypeError, ValueError) as error:
                    raise ValueError(f'{column} inválido en {design_path}: {row[column]!r}') from error
                if not math.isfinite(value) or value <= 0 or not value.is_integer():
                    raise ValueError(f'{column} debe contener enteros positivos en {design_path}')
                identifiers[column] = int(value)
            if row['tratamiento'] not in {'Pina', 'Bosque'}:
                raise ValueError(f'Tratamiento no reconocido en {design_path}')
            key = (row['tratamiento'], identifiers['lote'])
            if key in design:
                raise ValueError(f'Parcela duplicada en {design_path}: {key}')
            design[key] = identifiers['finca_id']
            producer = (row['productor_ficha'] or '').strip()
            if producer:
                farm = identifiers['finca_id']
                if farm in producer_by_farm and producer_by_farm[farm] != producer:
                    raise ValueError(f'Hay nombres de productor distintos para la finca {farm} en {design_path}')
                producer_by_farm[farm] = producer

    public = []
    for result in results:
        key = (result['tratamiento'], result['lote'])
        if key not in design:
            raise ValueError(f'Parcela sin correspondencia en diseno_muestreo.csv: {result["parcela"]}')
        record = {**result, 'finca_id': design[key],
                  'productor_finca': producer_by_farm.get(design[key], '')}
        public.append({column: record[column] for column in PUBLIC_COLUMNS})
    return public


def write_csv(path: Path, columns: list[str], records: list[dict], encoding: str = 'utf-8-sig') -> None:
    """Conserva los valores originales; None se exporta como celda vacía."""
    with path.open('w', encoding=encoding, newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(records)


def read_chemical_records(source: Path = SOURCE) -> list[dict]:
    book = load_workbook(source, data_only=True, read_only=True)
    try:
        return _read_sheet(book, source)
    finally:
        book.close()


def _read_sheet(book, source: Path) -> list[dict]:
    # El reporte incluye plantillas ocultas con cabeceras pero sin muestras.
    sheets = [sheet for sheet in book if sheet.sheet_state == 'visible'
              and sheet['A18'].value == 'ID USUARIO' and sheet['D18'].value == 'ID LAB']
    if len(sheets) != 1:
        raise ValueError('No se identifica una hoja química única con el formato esperado')
    sheet = sheets[0]
    # Validar el formato evita leer niveles críticos como observaciones.
    units = {'E16': 'pH', 'F16': 'cmol(+)/L', 'K16': '%', 'L16': 'mg/L',
             'Q16': 'mS/cm', 'R16': '%', 'T16': 'Relación'}
    for coordinate, unit in units.items():
        if sheet[coordinate].value != unit:
            raise ValueError(f'Unidad inesperada en {coordinate}: {sheet[coordinate].value!r}')
    for col, (_, header, _) in enumerate(VARIABLES, 5):
        if sheet.cell(17, col).value != header:
            raise ValueError(f'Encabezado inesperado en columna {col}')
    records = []
    for row in sheet.iter_rows(min_row=19):
        user_id, lab_id = row[0].value, row[3].value
        if not re.fullmatch(r'S-\d{2}-\d+', str(lab_id)):
            if re.match(r'(PIÑA|PINA|BOSQUE)\s+\d+', str(user_id), flags=re.I):
                raise ValueError(f'Muestra con ID de laboratorio inválido en fila {row[0].row}')
            continue
        match = re.fullmatch(r'(PIÑA|BOSQUE)\s+(\d+)\s*(?:\(.*\))?', str(user_id))
        if not match:
            raise ValueError(f'Parcela no reconocida: {user_id}')
        treatment = 'Pina' if match[1] == 'PIÑA' else 'Bosque'
        for col, (variable, _, unit) in enumerate(VARIABLES, 4):
            cell = row[col]
            number, operator, limit, state = parse_result(cell.value)
            records.append(dict(
                parcela=f'{treatment}_{int(match[2])}', tratamiento=treatment,
                lote=int(match[2]), id_usuario_original=user_id, id_lab=lab_id,
                variable=variable, unidad=unit, resultado_original=cell.value,
                valor=number, operador=operator, limite_reportado=limit,
                estado=state, archivo=source.relative_to(ROOT).as_posix() if source.is_relative_to(ROOT) else source.name,
                hoja=sheet.title, celda=cell.coordinate,
            ))
    if not records or len({(r['parcela'], r['variable']) for r in records}) != len(records) or len({(r['id_lab'], r['variable']) for r in records}) != len(records):
        raise ValueError('No hay resultados o existen muestras duplicadas')
    return records


def main():
    records = read_chemical_records()
    public = prepare_dashboard_data(records)
    columns = list(records[0])
    non_quantified = [record for record in records if record['estado'] != 'cuantificado']
    # Equivalente a la vista pivotada: una fila por parcela e ID LAB.
    variables = sorted(variable for variable, _, _ in VARIABLES)
    by_parcel = {}
    for record in records:
        key = (record['parcela'], record['id_lab'])
        row = by_parcel.setdefault(key, {'parcela': key[0], 'id_lab': key[1]})
        row[record['variable']] = record['resultado_original']
    parcel_rows = [by_parcel[key] for key in sorted(by_parcel)]

    OUT.mkdir(parents=True, exist_ok=True)
    write_csv(OUT / 'resultados_quimicos.csv', columns, records)
    write_csv(OUT / 'resultados_no_cuantificados.csv', columns, non_quantified)
    # Vista legible: conserva literalmente <1; para calcular usar el archivo largo.
    write_csv(OUT / 'quimica_por_parcela.csv', ['parcela', 'id_lab', *variables], parcel_rows)
    dashboard_path = OUT / 'web' / 'resultados_quimicos.csv'
    dashboard_path.parent.mkdir(parents=True, exist_ok=True)
    write_csv(dashboard_path, PUBLIC_COLUMNS, public, encoding='utf-8')
    print(f'{len({record["parcela"] for record in records})} parcelas, {len(records)} resultados, '
          f'{sum(record["estado"] == "censurado" for record in records)} censurados, '
          f'{sum(record["estado"] == "faltante" for record in records)} faltantes.')
    print(f'Exportación web: {dashboard_path.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
